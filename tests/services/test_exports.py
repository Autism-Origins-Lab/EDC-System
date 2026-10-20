from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.database.connection import get_connection
from app.database.repositories.forms import MULLEN_DOMAINS, save_mullen_assessment
from app.database.repositories.patients import create_patient
from app.services.exports import export_patient_workbook


def test_export_patient_workbook_creates_expected_sheets(temp_database, tmp_path):
    output_path = tmp_path / "exports" / "edc_export.xlsx"
    assert export_patient_workbook(output_path) == output_path

    workbook = load_workbook(output_path)
    try:
        assert workbook.sheetnames == [
            "Patients",
            "Telephone Screening",
            "Screening Questionnaire",
            "Medical History",
            "Family Medical History",
            "Procedure Schedule",
            "Mullen Assessment",
            "Mullen Scores",
        ]
        assert all(sheet.max_row == 1 for sheet in workbook)
        assert "patient_id" in [cell.value for cell in workbook["Mullen Assessment"][1]]
        assert "raw_score" in [cell.value for cell in workbook["Mullen Scores"][1]]
    finally:
        workbook.close()

    with get_connection() as connection:
        logs = connection.execute("SELECT output_path FROM export_logs").fetchall()
    assert [row["output_path"] for row in logs] == [str(output_path)]
    assert list(output_path.parent.iterdir()) == [output_path]


def test_export_preserves_mullen_context_scores_and_patient_links(temp_database, tmp_path):
    first_patient = create_patient("EXPORT-001")
    second_patient = create_patient("EXPORT-002")
    context = {
        "assessment_date": "2025-03-01",
        "scoring_date_of_birth": "2025-01-01",
        "confidence_level": 90,
        "age_calculation": "calendar",
    }
    scores = {
        domain: {
            "raw_score": 0,
            "t_score": 20,
            "band_of_error": 2.5,
            "percentile_rank": 1,
            "descriptive_category": "Very Low",
            "age_equivalence": None,
        }
        for domain in MULLEN_DOMAINS
    }
    save_mullen_assessment(first_patient, scores, context=context)
    second_scores = {
        "gross_motor": {
            "raw_score": 3,
            "t_score": 37,
            "band_of_error": 11,
            "percentile_rank": 10,
            "descriptive_category": "Below Average",
            "age_equivalence": 1,
        },
        "fine_motor": {"raw_score": None},
    }
    second_context = {**context, "confidence_level": 95, "age_calculation": "30_day_borrow"}
    save_mullen_assessment(second_patient, second_scores, context=second_context)

    output_path = export_patient_workbook(tmp_path / "mullen.xlsx")
    workbook = load_workbook(output_path)
    try:
        def records(sheet_name):
            rows = workbook[sheet_name].iter_rows(values_only=True)
            headers = next(rows)
            return [dict(zip(headers, row)) for row in rows]

        patients = {row["id"]: row for row in records("Patients")}
        assert patients[first_patient]["subject_id"] == "EXPORT-001"
        assert patients[second_patient]["subject_id"] == "EXPORT-002"
        assessments = {row["patient_id"]: row for row in records("Mullen Assessment")}
        assert set(assessments) == {first_patient, second_patient}
        for patient_id, expected in ((first_patient, context), (second_patient, second_context)):
            assert {field: assessments[patient_id][field] for field in expected} == expected

        exported_scores = records("Mullen Scores")
        assert len(exported_scores) == 7
        for patient_id, expected_scores in ((first_patient, scores), (second_patient, second_scores)):
            domains = {
                row["domain"]: row
                for row in exported_scores
                if row["mullen_assessments_id"] == assessments[patient_id]["id"]
            }
            assert set(domains) == set(expected_scores)
            for domain, expected in expected_scores.items():
                assert {field: domains[domain][field] for field in expected} == expected
        assert all(
            domains["fine_motor"][field] is None
            for field in (
                "raw_score", "t_score", "band_of_error", "percentile_rank",
                "descriptive_category", "age_equivalence",
            )
        )
    finally:
        workbook.close()


@pytest.mark.parametrize("existing_workbook", [False, True])
def test_failed_export_preserves_destination_and_does_not_log(
    temp_database, tmp_path, monkeypatch, existing_workbook
):
    output_path = tmp_path / "export.xlsx"
    original = None
    if existing_workbook:
        export_patient_workbook(output_path)
        original = output_path.read_bytes()

    read_sql = pd.read_sql_query

    def fail_on_scores(sql, connection):
        if sql == "SELECT * FROM mullen_scores":
            raise pd.errors.DatabaseError("Simulated score read failure")
        return read_sql(sql, connection)

    monkeypatch.setattr(pd, "read_sql_query", fail_on_scores)
    with pytest.raises(pd.errors.DatabaseError, match="Simulated score read failure"):
        export_patient_workbook(output_path)

    if existing_workbook:
        assert output_path.read_bytes() == original
    else:
        assert not output_path.exists()
    assert not list(tmp_path.glob(".edc-export-*"))
    with get_connection() as connection:
        count = connection.execute("SELECT COUNT(*) FROM export_logs").fetchone()[0]
    assert count == int(existing_workbook)


def test_locked_destination_does_not_log_success(temp_database, tmp_path, monkeypatch):
    output_path = export_patient_workbook(tmp_path / "export.xlsx")
    original = output_path.read_bytes()

    def deny_replace(self, target):
        raise PermissionError("Workbook is open in Excel")

    monkeypatch.setattr(Path, "replace", deny_replace)
    with pytest.raises(PermissionError, match="Workbook is open in Excel"):
        export_patient_workbook(output_path)

    assert output_path.read_bytes() == original
    assert not list(tmp_path.glob(".edc-export-*"))
    with get_connection() as connection:
        count = connection.execute("SELECT COUNT(*) FROM export_logs").fetchone()[0]
    assert count == 1
